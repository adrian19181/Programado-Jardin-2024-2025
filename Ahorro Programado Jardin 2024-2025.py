import sys
import io
import requests
import openpyxl
import pandas as pd
import numpy as np
import streamlit as st
import plotly.graph_objects as go

# ==============================================================================
# AUTORUNNER: FUNCIONA TANTO LOCALMENTE COMO EN STREAMLIT CLOUD
# ==============================================================================
if __name__ == "__main__":
    if not st.runtime.exists():
        from streamlit.web import cli as stcli
        sys.argv = ["streamlit", "run", sys.argv[0]]
        sys.exit(stcli.main())

# ==============================================================================
# CONFIGURACIÓN DE PÁGINA
# ==============================================================================
st.set_page_config(
    page_title="Ahorro Programado Jardín 2024-2025",
    page_icon="💰",
    layout="wide"
)

DRIVE_FILE_ID = "1ehwdQzNEeu3jhQWTlPhcezUMfCHtNute"

# ==============================================================================
# CARGAR DESDE GOOGLE DRIVE (CON CACHÉ REUTILIZABLE)
# ==============================================================================
@st.cache_data(ttl=300)
def cargar_excel_drive(file_id):
    url = f"https://drive.google.com/uc?id={file_id}&export=download"
    session = requests.Session()
    response = session.get(url)
    
    if response.status_code != 200:
        st.error("No se pudo descargar el archivo desde Google Drive.")
        return None

    excel_bytes = io.BytesIO(response.content)
    wb = openpyxl.load_workbook(excel_bytes, data_only=True)
    return wb

def obtener_valor(sheet, celda_principal, celda_secundaria=None):
    val = sheet[celda_principal].value
    if val is None and celda_secundaria:
        val = sheet[celda_secundaria].value
    return val

def fmt_moneda(val):
    if val is None:
        return "$0,00"
    try:
        f = float(val)
        return f"${f:,.2f}".replace(".", "X").replace(",", ".").replace("X", ",")
    except:
        return str(val)

def fmt_porcentaje(val):
    if val is None:
        return "0,00%"
    try:
        f = float(val)
        if f <= 1:
            f = f * 100
        return f"{f:.2f}%".replace(".", ",")
    except:
        return str(val)

# ==============================================================================
# ENCABEZADO PRINCIPAL Y RECARGA
# ==============================================================================
col_title, col_btn = st.columns([4, 1])
with col_title:
    st.title("💰 Ahorro Programado Jardín 2024-2025")
    st.caption("Sincronizado en tiempo real desde Google Drive")

with col_btn:
    if st.button("🔄 Recargar", use_container_width=True):
        st.cache_data.clear()
        st.rerun()

st.divider()

wb = cargar_excel_drive(DRIVE_FILE_ID)

if wb is not None:
    sheet = wb.active

    # ==============================================================================
    # 1. FICHA TÉCNICA CONDICIONES DEL AHORRO (E1:J9)
    # ==============================================================================
    tna = obtener_valor(sheet, "H1", "F1")
    tasa_diaria = obtener_valor(sheet, "J2", "F2")
    fecha_inicio = obtener_valor(sheet, "H3", "F3")
    plazo_meses = obtener_valor(sheet, "H4", "F4")
    fecha_fin = obtener_valor(sheet, "H5", "F5")
    total_dias = obtener_valor(sheet, "H6", "F6")
    valor_aporte = obtener_valor(sheet, "H7", "F7")
    dia_aporte = obtener_valor(sheet, "H8", "F8")
    meses_real = sheet["H9"].value
    dias_real = sheet["J9"].value

    tna_num = float(tna) if isinstance(tna, (int, float)) else 0.065
    tna_val = tna_num * 100 if tna_num <= 1 else tna_num
    fecha_ini_str = fecha_inicio.strftime('%d/%m/%Y') if hasattr(fecha_inicio, 'strftime') else str(fecha_inicio or "-")
    fecha_fin_str = fecha_fin.strftime('%d/%m/%Y') if hasattr(fecha_fin, 'strftime') else str(fecha_fin or "-")
    val_diaria_str = f"{float(tasa_diaria):.7f}".replace(".", ",") if isinstance(tasa_diaria, (int, float)) else str(tasa_diaria)
    val_aporte_str = f"{float(valor_aporte):,.2f}".replace(".", ",") if isinstance(valor_aporte, (int, float)) else str(valor_aporte)

    st.subheader("📋 Condiciones del Ahorro Programado")

    html_condiciones = f"""<style>
.excel-tbl-card {{
    background-color: #ffffff;
    padding: 10px;
    border-radius: 6px;
    box-shadow: 0 3px 10px rgba(0, 0, 0, 0.2);
    display: inline-block;
    margin-bottom: 15px;
}}
.excel-tbl {{
    border-collapse: collapse;
    font-family: Calibri, 'Segoe UI', Arial, sans-serif;
    font-size: 13px;
    color: #000000;
    width: auto;
}}
.excel-tbl td {{
    border: 1px solid #000000;
    padding: 4px 8px;
    vertical-align: middle;
}}
.lbl-yellow {{
    background-color: #FFE600;
    font-weight: bold;
    text-align: right;
    white-space: nowrap;
}}
.val-green {{
    background-color: #E2EFDA;
    text-align: center;
    font-weight: 600;
}}
.val-formula {{
    background-color: #D9E1F2;
    text-align: center;
    font-size: 12px;
}}
.val-tan {{
    background-color: #FFF2CC;
    text-align: center;
    font-weight: 600;
}}
</style>
<div class="excel-tbl-card">
<table class="excel-tbl">
<tr>
<td class="lbl-yellow">Tasa Nominal Anual</td>
<td colspan="4" class="val-green">{tna_val:.2f}%</td>
</tr>
<tr>
<td class="lbl-yellow">Tasa Diaria</td>
<td colspan="3" class="val-formula">((TNA/100)/360)</td>
<td class="val-green" style="font-weight: bold;">{val_diaria_str}</td>
</tr>
<tr>
<td class="lbl-yellow">Fecha inicio</td>
<td colspan="4" class="val-green">{fecha_ini_str}</td>
</tr>
<tr>
<td class="lbl-yellow">Plazo meses</td>
<td colspan="4" class="val-green">{plazo_meses}</td>
</tr>
<tr>
<td class="lbl-yellow">Fecha Fin</td>
<td colspan="4" class="val-tan">{fecha_fin_str}</td>
</tr>
<tr>
<td class="lbl-yellow">Total de Días</td>
<td colspan="4" class="val-tan">{total_dias}</td>
</tr>
<tr>
<td class="lbl-yellow">Valor Aporte Mensual</td>
<td colspan="4" class="val-green">${val_aporte_str}</td>
</tr>
<tr>
<td class="lbl-yellow">Dia de Aporte Mensual</td>
<td colspan="4" class="val-green">{dia_aporte}</td>
</tr>
<tr>
<td class="lbl-yellow">Cantidad de Meses Real</td>
<td class="val-tan" style="width: 15%;">Meses</td>
<td class="val-tan" style="width: 15%;">{meses_real}</td>
<td class="val-tan" style="width: 15%;">Días</td>
<td class="val-tan" style="width: 15%;">{dias_real}</td>
</tr>
</table>
</div>"""

    st.markdown(html_condiciones, unsafe_allow_html=True)

    st.divider()

    # ==============================================================================
    # 2. SECCIONES DE KPIS: RESULTADOS Y RENDIMIENTOS
    # ==============================================================================
    d2 = sheet["D2"].value   # Total Interés Ganados
    d3 = sheet["D3"].value   # Total Depósitos Mensuales
    d4 = sheet["D4"].value   # Total Depósitos Personales
    d5 = sheet["D5"].value   # Total Depositado sin intereses
    d6 = sheet["D6"].value   # Total Ganado Incluido Intereses

    b7 = sheet["B7"].value   # TIR (Tasa Interna de Retorno)
    d7 = sheet["D7"].value   # ROI RENTABILIDAD ACUMULADA
    d16 = sheet["D16"].value # Saldo Fin mes Promedio Ponderado
    d19 = sheet["D19"].value # Aportes Personales Promedio
    b19 = sheet["B19"].value # Interés Diario Promedio

    # CÁLCULOS DE KPIS
    tna_dec = tna_num if tna_num <= 1 else tna_num / 100.0
    tea_calc = ((1.0 + (tna_dec / 12.0)) ** 12.0) - 1.0

    d2_num = float(d2) if isinstance(d2, (int, float)) else 0.0
    d6_num = float(d6) if isinstance(d6, (int, float)) else 0.0
    ratio_ganancia_calc = (d2_num / d6_num) if d6_num > 0 else 0.0

    css_kpis = """<style>
.kpi-grid {
    display: flex;
    flex-wrap: wrap;
    gap: 10px;
    margin-top: 8px;
    margin-bottom: 20px;
}
.kpi-card {
    background-color: #ffffff;
    border: 1px solid #000000;
    border-radius: 6px;
    padding: 8px 10px;
    flex: 1 1 140px;
    min-width: 125px;
    box-shadow: 0 2px 6px rgba(0,0,0,0.15);
    text-align: center;
    font-family: Calibri, 'Segoe UI', Arial, sans-serif;
}
.kpi-card-yellow {
    border-top: 4px solid #FFE600;
    background-color: #FFFFF0;
}
.kpi-card-green {
    border-top: 4px solid #22C55E;
    background-color: #F0FDF4;
}
.kpi-card-blue {
    border-top: 4px solid #3B82F6;
    background-color: #EFF6FF;
}
.kpi-card-purple {
    border-top: 4px solid #8B5CF6;
    background-color: #F5F3FF;
}
.kpi-label {
    font-size: 11px;
    font-weight: bold;
    color: #374151;
    text-transform: uppercase;
    line-height: 1.2;
    margin-bottom: 4px;
}
.kpi-value {
    font-size: 16px;
    font-weight: bold;
    color: #000000;
}
</style>"""

    # --- KPI RESULTADOS ---
    st.subheader("📊 KPI Resultados")
    
    html_kpi_resultados = f"""{css_kpis}
<div class="kpi-grid">
<div class="kpi-card kpi-card-green">
<div class="kpi-label">Total Interés Ganados</div>
<div class="kpi-value">{fmt_moneda(d2)}</div>
</div>
<div class="kpi-card kpi-card-blue">
<div class="kpi-label">Total Dep. Mensuales</div>
<div class="kpi-value">{fmt_moneda(d3)}</div>
</div>
<div class="kpi-card kpi-card-yellow">
<div class="kpi-label">Total Dep. Personales</div>
<div class="kpi-value">{fmt_moneda(d4)}</div>
</div>
<div class="kpi-card kpi-card-blue">
<div class="kpi-label">Total Depositado (sin Int.)</div>
<div class="kpi-value">{fmt_moneda(d5)}</div>
</div>
<div class="kpi-card kpi-card-green">
<div class="kpi-label">Total Ganado Incluido Intereses</div>
<div class="kpi-value">{fmt_moneda(d6)}</div>
</div>
</div>"""

    st.markdown(html_kpi_resultados, unsafe_allow_html=True)

    # --- KPI RENDIMIENTOS ---
    st.subheader("📈 KPI Rendimientos")

    html_kpi_rendimientos = f"""<div class="kpi-grid">
<div class="kpi-card kpi-card-purple">
<div class="kpi-label">TIR (Tasa Int. Retorno)</div>
<div class="kpi-value">{fmt_porcentaje(b7)}</div>
</div>
<div class="kpi-card kpi-card-purple">
<div class="kpi-label">TEA (Tasa Efe. Anual)</div>
<div class="kpi-value">{fmt_porcentaje(tea_calc)}</div>
</div>
<div class="kpi-card kpi-card-green">
<div class="kpi-label">ROI Rentabilidad Acum.</div>
<div class="kpi-value">{fmt_porcentaje(d7)}</div>
</div>
<div class="kpi-card kpi-card-purple">
<div class="kpi-label">Ratio de Ganancia</div>
<div class="kpi-value">{fmt_porcentaje(ratio_ganancia_calc)}</div>
</div>
<div class="kpi-card kpi-card-blue">
<div class="kpi-label">Saldo Fin Mes Prom. Pond.</div>
<div class="kpi-value">{fmt_moneda(d16)}</div>
</div>
<div class="kpi-card kpi-card-yellow">
<div class="kpi-label">Aportes Personales Prom.</div>
<div class="kpi-value">{fmt_moneda(d19)}</div>
</div>
<div class="kpi-card kpi-card-green">
<div class="kpi-label">Interés Diario Promedio</div>
<div class="kpi-value">{fmt_moneda(b19)}</div>
</div>
</div>"""

    st.markdown(html_kpi_rendimientos, unsafe_allow_html=True)

    st.divider()

    # ==============================================================================
    # 3. TABLA RESUMEN DE MOVIMIENTOS DINÁMICA (AZ:BF)
    # ==============================================================================
    st.subheader("📑 Tabla Resumen de Movimientos")

    movements = []
    r = 3  # Fila de inicio (AZ3:BF3)
    
    while True:
        date_val = sheet.cell(row=r, column=52).value  # Col AZ (Fechas)
        if date_val is None or str(date_val).strip() == "":
            break

        dep_m = sheet.cell(row=r, column=53).value or 0     # Col BA
        dep_p = sheet.cell(row=r, column=54).value or 0     # Col BB
        saldo_acum = sheet.cell(row=r, column=55).value or 0 # Col BC
        int_diario = sheet.cell(row=r, column=56).value or 0 # Col BD
        inc_int_d = sheet.cell(row=r, column=57).value or 0  # Col BE
        int_m = sheet.cell(row=r, column=58).value          # Col BF

        is_int_m = int_m is not None and isinstance(int_m, (int, float)) and float(int_m) > 0
        is_dep_p = isinstance(dep_p, (int, float)) and float(dep_p) > 0
        is_dep_m = isinstance(dep_m, (int, float)) and float(dep_m) > 0

        if is_int_m:
            bg_color = "#B4C6E7"  # Azul suave
            tipo_cat = "🟦 Intereses Ganados"
        elif is_dep_p:
            bg_color = "#FFFF00"  # Amarillo
            tipo_cat = "🟨 Aportes Personales"
        elif is_dep_m:
            bg_color = "#FCE4D6"  # Naranja/Rosado
            tipo_cat = "🟧 Aportes Mensuales"
        else:
            bg_color = "#FFFFFF"
            tipo_cat = "Otros"

        movements.append({
            'fecha': date_val,
            'dep_m': float(dep_m) if isinstance(dep_m, (int, float)) else 0.0,
            'dep_p': float(dep_p) if isinstance(dep_p, (int, float)) else 0.0,
            'saldo_acum': float(saldo_acum) if isinstance(saldo_acum, (int, float)) else 0.0,
            'int_diario': float(int_diario) if isinstance(int_diario, (int, float)) else 0.0,
            'inc_int_d': float(inc_int_d) if isinstance(inc_int_d, (int, float)) else 0.0,
            'int_m': float(int_m) if is_int_m else None,
            'bg_color': bg_color,
            'tipo_cat': tipo_cat
        })
        r += 1

    opciones_disponibles = ["🟨 Aportes Personales", "🟧 Aportes Mensuales", "🟦 Intereses Ganados"]
    
    st.write("**Filtra los movimientos por categoría:**")
    
    if hasattr(st, "pills"):
        categorias_seleccionadas = st.pills(
            label="Categorías",
            options=opciones_disponibles,
            default=opciones_disponibles,
            selection_mode="multi",
            label_visibility="collapsed"
        )
    else:
        categorias_seleccionadas = st.multiselect(
            label="Selecciona Categorías a mostrar",
            options=opciones_disponibles,
            default=opciones_disponibles,
            label_visibility="collapsed"
        )

    movimientos_filtrados = [
        m for m in movements if m['tipo_cat'] in (categorias_seleccionadas or [])
    ]

    rows_html = []
    for m in movimientos_filtrados:
        dt_str = m['fecha'].strftime('%d/%m/%Y') if hasattr(m['fecha'], 'strftime') else str(m['fecha'])
        dep_m_str = f"{m['dep_m']:,.2f}".replace(".", "X").replace(",", ".").replace("X", ",") if m['dep_m'] > 0 else "0,00"
        dep_p_str = f"{m['dep_p']:,.2f}".replace(".", "X").replace(",", ".").replace("X", ",") if m['dep_p'] > 0 else "0,00"
        saldo_str = f"{m['saldo_acum']:,.2f}".replace(".", "X").replace(",", ".").replace("X", ",")
        int_d_str = f"{m['int_diario']:.2f}".replace(".", ",")
        inc_d_str = f"{m['inc_int_d']:.2f}".replace(".", ",") if m['inc_int_d'] > 0 else ""
        int_m_str = f"<b>{m['int_m']:,.2f}</b>".replace(".", "X").replace(",", ".").replace("X", ",") if m['int_m'] is not None else ""

        rows_html.append(
            f'<tr style="background-color: {m["bg_color"]};">'
            f'<td><b>{dt_str}</b></td>'
            f'<td>{dep_m_str}</td>'
            f'<td>{dep_p_str}</td>'
            f'<td style="font-weight: 600;">{saldo_str}</td>'
            f'<td>{int_d_str}</td>'
            f'<td>{inc_d_str}</td>'
            f'<td>{int_m_str}</td>'
            f'</tr>'
        )

    tabla_movs_body = "".join(rows_html)

    html_movs_completo = f"""<style>
.tbl-scroll-wrapper {{
    max-height: 480px;
    max-width: 100%;
    overflow-y: auto;
    overflow-x: auto;
    border: 1px solid #000000;
    border-radius: 6px;
    background-color: #ffffff;
    box-shadow: 0 4px 12px rgba(0,0,0,0.2);
    margin-top: 10px;
    margin-bottom: 25px;
    display: inline-block;
}}
.tbl-sticky-movs {{
    width: auto;
    border-collapse: collapse;
    font-family: Calibri, 'Segoe UI', Arial, sans-serif;
    font-size: 12px;
    color: #000000;
}}
.tbl-sticky-movs th {{
    position: sticky;
    top: 0;
    z-index: 10;
    background-color: #1F2937;
    color: #ffffff;
    border: 1px solid #000000;
    padding: 5px 6px;
    text-align: center;
    font-weight: bold;
    line-height: 1.15;
    white-space: nowrap;
}}
.tbl-sticky-movs td {{
    border: 1px solid #000000;
    padding: 4px 6px;
    vertical-align: middle;
    text-align: center !important;
    white-space: nowrap;
}}
</style>
<div class="tbl-scroll-wrapper">
<table class="tbl-sticky-movs">
<thead>
<tr>
<th>Fechas</th>
<th>Dep.<br>Mensual</th>
<th>Dep.<br>Personales</th>
<th>Saldo<br>Acumulado</th>
<th>Interés<br>Diario</th>
<th>Incr. Int. Diario<br>por Aportes</th>
<th>Interés<br>Mensual</th>
</tr>
</thead>
<tbody>
{tabla_movs_body if rows_html else '<tr><td colspan="7" style="text-align:center; padding:15px; font-weight:bold;">No hay movimientos para los filtros seleccionados</td></tr>'}
</tbody>
</table>
</div>"""

    st.markdown(html_movs_completo, unsafe_allow_html=True)

    st.divider()

    # ==============================================================================
    # 4. SECCIÓN DE GRÁFICOS HORIZONTALES INTELIGENTES 100% ADAPTADOS A MÓVIL
    # ==============================================================================
    st.subheader("📊 Evolución de Saldo e Intereses Mensuales")

    # 4.1 LECTURA DE RESUMEN MENSUAL (AC:AG)
    chart_dates = []
    chart_saldos = []
    chart_intereses = []

    r_chart = 2 # Fila inicio en tabla resumen mensual (AC2:AG14)
    while True:
        f_cobro = sheet.cell(row=r_chart, column=29).value # Col AC (Fecha Cobro Intereses)
        saldo_fin = sheet.cell(row=r_chart, column=30).value # Col AD (Saldo Fin mes)
        int_mes = sheet.cell(row=r_chart, column=33).value # Col AG (Interés Mensual)

        if f_cobro is None or str(f_cobro).strip() == "":
            break

        dt_str = f_cobro.strftime('%d/%m/%Y') if hasattr(f_cobro, 'strftime') else str(f_cobro)
        saldo_val = float(saldo_fin) if isinstance(saldo_fin, (int, float)) else 0.0
        int_mes_val = float(int_mes) if isinstance(int_mes, (int, float)) else 0.0

        chart_dates.append(dt_str)
        chart_saldos.append(saldo_val)
        chart_intereses.append(int_mes_val)
        r_chart += 1

    # 4.2 CÁLCULO DE INTERÉS DIARIO PROMEDIO POR MES (COL A Y R CON FALLBACK)
    daily_int_by_month = {}
    
    # Intento 1: Leer directamente desde Columna A (1) y Columna R (18)
    for r_d in range(2, sheet.max_row + 1):
        f_val = sheet.cell(row=r_d, column=1).value  # Col A (Fechas)
        r_val = sheet.cell(row=r_d, column=18).value # Col R (Interés Diario)

        if f_val is None or str(f_val).strip() == "":
            continue

        dt = None
        if hasattr(f_val, 'year'):
            dt = f_val
        else:
            try:
                dt = pd.to_datetime(f_val, dayfirst=True)
            except:
                continue

        try:
            num_val = float(r_val) if r_val is not None else 0.0
        except (ValueError, TypeError):
            continue

        if num_val > 0:
            m_key = (dt.year, dt.month)
            if m_key not in daily_int_by_month:
                daily_int_by_month[m_key] = []
            daily_int_by_month[m_key].append(num_val)

    # Intento 2 (Fallback): Si las fórmulas en Col A/R no tienen valor en caché, usas movimientos
    if not daily_int_by_month and movements:
        for m in movements:
            f_val = m['fecha']
            r_val = m['int_diario']
            if f_val is not None and r_val > 0:
                dt = f_val if hasattr(f_val, 'year') else pd.to_datetime(f_val, dayfirst=True)
                m_key = (dt.year, dt.month)
                if m_key not in daily_int_by_month:
                    daily_int_by_month[m_key] = []
                daily_int_by_month[m_key].append(r_val)

    avg_daily_dates = []
    avg_daily_vals = []

    for (year, month), vals in sorted(daily_int_by_month.items()):
        start_month_str = f"01/{month:02d}/{year}"
        avg_val = sum(vals) / len(vals) if vals else 0.0
        avg_daily_dates.append(start_month_str)
        avg_daily_vals.append(avg_val)

    # 4.3 CREACIÓN DE PESTAÑAS Y GRÁFICOS
    tab_int, tab_saldo, tab_avg, tab_cum = st.tabs([
        "🔴 Interés Ganado Mensual", 
        "🟦 Saldo Fin de Mes", 
        "🟢 Interés Diario Promedio",
        "📈 Interés Acumulado Creciente"
    ])

    # PESTAÑA 1: INTERÉS GANADO MENSUAL
    with tab_int:
        if chart_intereses:
            max_int = max(chart_intereses) if chart_intereses else 1.0
            pos_int = ['outside' if v < max_int * 0.35 else 'inside' for v in chart_intereses]
            
            fig_int = go.Figure()
            fig_int.add_trace(go.Bar(
                x=chart_intereses,
                y=chart_dates,
                orientation='h',
                marker_color='#EF4444',
                text=[f"${v:,.2f}".replace(".", "X").replace(",", ".").replace("X", ",") for v in chart_intereses],
                textposition=pos_int,
                textfont=dict(size=11),
                insidetextfont=dict(color='#FFFFFF'),
                outsidetextfont=dict(color='#FAFAFA'),
                hoverinfo='none'
            ))
            fig_int.update_layout(
                template="plotly_dark",
                margin=dict(l=75, r=65, t=10, b=30),
                height=max(380, len(chart_dates) * 32),
                xaxis=dict(showgrid=True, gridcolor='#334155', title=""),
                yaxis=dict(autorange="reversed", tickfont=dict(size=10, color='#CBD5E1')),
                paper_bgcolor="rgba(0,0,0,0)",
                plot_bgcolor="rgba(0,0,0,0)"
            )
            st.plotly_chart(
                fig_int, 
                use_container_width=True, 
                config={'staticPlot': True, 'displayModeBar': False}
            )

    # PESTAÑA 2: SALDO FIN DE MES
    with tab_saldo:
        if chart_saldos:
            max_sal = max(chart_saldos) if chart_saldos else 1.0
            pos_sal = ['outside' if v < max_sal * 0.35 else 'inside' for v in chart_saldos]

            fig_sal = go.Figure()
            fig_sal.add_trace(go.Bar(
                x=chart_saldos,
                y=chart_dates,
                orientation='h',
                marker_color='#3B82F6',
                text=[f"${v:,.2f}".replace(".", "X").replace(",", ".").replace("X", ",") for v in chart_saldos],
                textposition=pos_sal,
                textfont=dict(size=10),
                insidetextfont=dict(color='#FFFFFF'),
                outsidetextfont=dict(color='#FAFAFA'),
                hoverinfo='none'
            ))
            fig_sal.update_layout(
                template="plotly_dark",
                margin=dict(l=75, r=65, t=10, b=30),
                height=max(380, len(chart_dates) * 32),
                xaxis=dict(showgrid=True, gridcolor='#334155', title=""),
                yaxis=dict(autorange="reversed", tickfont=dict(size=10, color='#CBD5E1')),
                paper_bgcolor="rgba(0,0,0,0)",
                plot_bgcolor="rgba(0,0,0,0)"
            )
            st.plotly_chart(
                fig_sal, 
                use_container_width=True, 
                config={'staticPlot': True, 'displayModeBar': False}
            )

    # PESTAÑA 3: INTERÉS DIARIO PROMEDIO
    with tab_avg:
        if avg_daily_vals:
            max_avg = max(avg_daily_vals) if avg_daily_vals else 1.0
            pos_avg = ['outside' if v < max_avg * 0.35 else 'inside' for v in avg_daily_vals]

            fig_avg = go.Figure()
            fig_avg.add_trace(go.Bar(
                x=avg_daily_vals,
                y=avg_daily_dates,
                orientation='h',
                marker_color='#22C55E',
                text=[f"${v:,.2f}".replace(".", "X").replace(",", ".").replace("X", ",") for v in avg_daily_vals],
                textposition=pos_avg,
                textfont=dict(size=11),
                insidetextfont=dict(color='#FFFFFF'),
                outsidetextfont=dict(color='#FAFAFA'),
                hoverinfo='none'
            ))
            fig_avg.update_layout(
                template="plotly_dark",
                margin=dict(l=75, r=65, t=10, b=30),
                height=max(380, len(avg_daily_dates) * 32),
                xaxis=dict(showgrid=True, gridcolor='#334155', title=""),
                yaxis=dict(autorange="reversed", tickfont=dict(size=10, color='#CBD5E1')),
                paper_bgcolor="rgba(0,0,0,0)",
                plot_bgcolor="rgba(0,0,0,0)"
            )
            st.plotly_chart(
                fig_avg, 
                use_container_width=True, 
                config={'staticPlot': True, 'displayModeBar': False}
            )

    # PESTAÑA 4: INTERÉS ACUMULADO CRECIENTE (SOLUCIÓN DE ESCALA)
    with tab_cum:
        if chart_saldos and chart_intereses:
            cum_int = np.cumsum(chart_intereses).tolist()
            pct_roi = [(i / s) * 100 for i, s in zip(cum_int, chart_saldos)]

            max_cum = max(cum_int) if cum_int else 1.0
            pos_cum = ['outside' if v < max_cum * 0.35 else 'inside' for v in cum_int]

            labels_cum = [
                f"${v:,.2f} ({p:.1f}%)".replace(".", "X").replace(",", ".").replace("X", ",")
                for v, p in zip(cum_int, pct_roi)
            ]

            fig_cum = go.Figure()
            fig_cum.add_trace(go.Bar(
                x=cum_int,
                y=chart_dates,
                orientation='h',
                marker_color='#10B981', # Verde Esmeralda
                text=labels_cum,
                textposition=pos_cum,
                textfont=dict(size=11),
                insidetextfont=dict(color='#FFFFFF'),
                outsidetextfont=dict(color='#FAFAFA'),
                hoverinfo='none'
            ))
            fig_cum.update_layout(
                template="plotly_dark",
                margin=dict(l=75, r=85, t=10, b=30),
                height=max(380, len(chart_dates) * 32),
                xaxis=dict(showgrid=True, gridcolor='#334155', title=""),
                yaxis=dict(autorange="reversed", tickfont=dict(size=10, color='#CBD5E1')),
                paper_bgcolor="rgba(0,0,0,0)",
                plot_bgcolor="rgba(0,0,0,0)"
            )
            st.plotly_chart(
                fig_cum, 
                use_container_width=True, 
                config={'staticPlot': True, 'displayModeBar': False}
            )